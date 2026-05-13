import json
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from pypdf import PdfReader
from pptx import Presentation


BASE_FILES_DIR = Path(__file__).parent.parent / "files"
BASE_FILES_DIR_RESOLVED = BASE_FILES_DIR.resolve()

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
    ".json",
}


def extract_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    text = []
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            text.append(extracted)
    return "\n".join(text)


def extract_docx(file_path: Path) -> str:
    doc = Document(str(file_path))
    return "\n".join([p.text for p in doc.paragraphs if p.text])


def extract_pptx(file_path: Path) -> str:
    prs = Presentation(str(file_path))
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_runs.append(shape.text)
    return "\n".join(text_runs)


def extract_txt_md(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8", errors="ignore")


def extract_json(file_path: Path) -> str:
    data = json.loads(file_path.read_text(encoding="utf-8"))
    return json.dumps(data, indent=2, ensure_ascii=False)


EXTRACTOR_REGISTRY = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".txt": extract_txt_md,
    ".md": extract_txt_md,
    ".json": extract_json,
}


def extract_file_content(file_path: Path) -> Dict[str, Any] | None:
    suffix = file_path.suffix.lower()
    extractor = EXTRACTOR_REGISTRY.get(suffix)

    base_result = {
        "name": file_path.name,
        "type": suffix.lstrip("."),
    }

    if not extractor:
        return None

    try:
        content = extractor(file_path)
        return {
            **base_result,
            "content": content.strip() if content else "",
        }
    except Exception as error:
        return {
            **base_result,
            "content": "",
            "error": str(error),
        }


def recursive_file_search(base_path: Path) -> List[Path]:
    return [
        file
        for file in base_path.rglob("*")
        if file.is_file() and file.suffix.lower() in SUPPORTED_EXTENSIONS
    ]


def execute(source_folder: Path | None = None) -> Dict[str, Any]:
    base_path = source_folder or BASE_FILES_DIR_RESOLVED
    files = recursive_file_search(base_path)

    list_result = [
        {
            "name": file.name,
            "type": file.suffix.lower().replace(".", ""),
        }
        for file in files
    ]

    extraction_result = [
        result
        for file in files
        if (result := extract_file_content(file)) is not None
    ]

    return {
        "status": "success",
        "total_files_found": len(files),
        "list_result": list_result,
        "extraction_result": extraction_result,
    }
